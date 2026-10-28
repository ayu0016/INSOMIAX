from __future__ import annotations

import os
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from fastapi import UploadFile
from pypdf import PdfReader

from ..errors import RadixError
from .models import ExtractedDocument

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SUPPORTED = {".pdf": PDF_MIME, ".docx": DOCX_MIME}


def configured_max_upload_bytes() -> int:
    try:
        max_bytes = int(os.getenv("RADIX_JD_MAX_UPLOAD_BYTES", str(10 * 1024 * 1024)))
    except ValueError:
        raise RadixError(500, "INVALID_SERVER_CONFIGURATION", "The configured upload limit is invalid.") from None
    if max_bytes <= 0:
        raise RadixError(500, "INVALID_SERVER_CONFIGURATION", "The configured upload limit must be positive.")
    return max_bytes


def _extract_pdf(data: bytes) -> str:
    if not data.startswith(b"%PDF-"):
        raise ValueError("Invalid PDF signature")
    reader = PdfReader(BytesIO(data), strict=True)
    if reader.is_encrypted:
        raise ValueError("Encrypted PDFs are not supported")
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _extract_docx(data: bytes) -> str:
    with ZipFile(BytesIO(data)) as archive:
        names = set(archive.namelist())
        if "[Content_Types].xml" not in names or "word/document.xml" not in names:
            raise ValueError("Invalid DOCX package")
    document = Document(BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join(part for part in parts if part.strip()).strip()


async def extract_upload(file: UploadFile, max_bytes: int) -> ExtractedDocument:
    filename = (file.filename or "").strip()
    extension = "." + filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    if extension not in SUPPORTED or file.content_type != SUPPORTED[extension]:
        raise RadixError(415, "UNSUPPORTED_FILE_TYPE", "Upload a PDF or DOCX file with a matching content type.", {"allowed_extensions": [".pdf", ".docx"]})
    data = await file.read(max_bytes + 1)
    if len(data) > max_bytes:
        raise RadixError(413, "UPLOAD_TOO_LARGE", "The uploaded document exceeds the configured size limit.", {"maximum_bytes": max_bytes})
    if not data:
        raise RadixError(422, "EMPTY_DOCUMENT", "The uploaded document is empty.")
    try:
        text = _extract_pdf(data) if extension == ".pdf" else _extract_docx(data)
    except Exception:
        # Parser exceptions vary between malformed files and dependency versions; never return their contents.
        raise RadixError(422, "UNREADABLE_DOCUMENT", "The document could not be parsed. Check that it is a valid, unencrypted PDF or DOCX.") from None
    if not text.strip():
        raise RadixError(422, "EMPTY_DOCUMENT_TEXT", "No readable text was found. Scanned PDFs require OCR, which is not available.")
    return ExtractedDocument(filename=filename, media_type=SUPPORTED[extension], text=text, character_count=len(text))
